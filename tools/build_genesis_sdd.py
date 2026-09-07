from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / "tools" / "genesis_reference.docx"
OUT = ROOT / "docs" / "GENESIS_Complete_System_Design_Specification_v1.0.docx"
SOURCE = ROOT / "docs" / "GENESIS_Complete_System_Design_Specification.md"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(95, 105, 115)


def set_font(style, name, size, color=None, bold=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def make_reference():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    set_font(normal, "Calibri", 11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Title", 25, DARK, 0, 12),
        ("Subtitle", 13, MUTED, 0, 16),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK, 10, 5),
        ("Heading 4", 11, DARK, 8, 4),
    ]:
        s = doc.styles[name]
        set_font(s, "Calibri", size, color, True if name != "Subtitle" else False)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        s = doc.styles[name]
        set_font(s, "Calibri", 11)
        s.paragraph_format.left_indent = Inches(0.375)
        s.paragraph_format.first_line_indent = Inches(-0.188)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.line_spacing = 1.25
    header = sec.header.paragraphs[0]
    header.text = "GENESIS | Complete System Design Specification"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.style, "Calibri", 9, MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    set_font(footer.style, "Calibri", 9, MUTED)
    doc.save(REF)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tc_mar.find(qn(f"w:{m}")) or OxmlElement(f"w:{m}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        if el.getparent() is None:
            tc_mar.append(el)


def postprocess():
    doc = Document(OUT)
    for sec in doc.sections:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1)
        sec.header_distance = sec.footer_distance = Inches(0.492)
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        cols = max(len(r.cells) for r in table.rows)
        if cols == 1:
            continue
        for ri, row in enumerate(table.rows):
            for _ci, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                set_cell_margins(cell)
                if ri == 0:
                    shade(cell, "E8EEF5")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.bold = True
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.1
                    for run in p.runs:
                        run.font.name = "Calibri"
                        run.font.size = Pt(9)
    doc.core_properties.title = "GENESIS Complete System Design Specification"
    doc.core_properties.subject = (
        "Implementation-ready specification for a local generative "
        "agent-based modelling research system"
    )
    doc.save(OUT)


def table_widths(headers):
    n = len(headers)
    if n == 2:
        return [1.1, 5.4]
    if n == 3 and headers[0].strip() == "ID":
        return [1.0, 4.0, 1.5]
    if n == 3:
        return [1.4, 2.55, 2.55]
    return [6.5 / n] * n


def add_fixed_table(doc, rows):
    headers = rows[0]
    widths = table_widths(headers)
    table = doc.add_table(rows=len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    dxa_widths = [round(w * 1440) for w in widths]
    dxa_widths[-1] += 9360 - sum(dxa_widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_pr.append(tbl_layout)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in dxa_widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for ri, values in enumerate(rows):
        if ri == 0:
            tr_pr = table.rows[ri]._tr.get_or_add_trPr()
            header_flag = OxmlElement("w:tblHeader")
            header_flag.set(qn("w:val"), "true")
            tr_pr.append(header_flag)
        for ci, value in enumerate(values):
            cell = table.cell(ri, ci)
            cell.width = Inches(dxa_widths[ci] / 1440)
            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER if ci == 0 else WD_CELL_VERTICAL_ALIGNMENT.TOP
            )
            set_cell_margins(cell)
            if ri == 0:
                shade(cell, "E8EEF5")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(value)
            run.bold = ri == 0
            run.font.name = "Calibri"
            run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_direct():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    set_font(normal, "Calibri", 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in [
        ("Title", 27, DARK, 0, 10),
        ("Subtitle", 13, MUTED, 0, 16),
        ("Heading 1", 16, BLUE, 18, 8),
        ("Heading 2", 13, BLUE, 14, 6),
        ("Heading 3", 11.5, DARK, 10, 4),
    ]:
        s = doc.styles[name]
        set_font(s, "Calibri", size, color, name != "Subtitle")
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        s = doc.styles[name]
        set_font(s, "Calibri", 10.5)
        s.paragraph_format.space_after = Pt(3)
        s.paragraph_format.line_spacing = 1.15
    code = doc.styles["No Spacing"]
    set_font(code, "Consolas", 8.5)
    header = sec.header.paragraphs[0]
    header.text = "GENESIS | Complete System Design Specification"
    set_font(header.style, "Calibri", 9, MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    set_font(footer.style, "Calibri", 9, MUTED)

    p = doc.add_paragraph(style="Title")
    p.add_run("GENESIS Complete System\nDesign Specification")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Complete local-first research system")
    for label, value in [
        ("Status", "Approved design baseline"),
        ("Version", "1.0"),
        ("Date", "30 August 2026"),
        ("Audience", "Researchers, research software engineers, reviewers, and coding agents"),
    ]:
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(value)
    doc.add_paragraph(
        "A concrete, implementation-ready specification derived from the "
        "GENESIS research architecture and source documents."
    )
    doc.add_page_break()

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings = [ln[3:].strip() for ln in lines if ln.startswith("## ")]
    doc.add_heading("Contents", level=1)
    for h in headings:
        p = doc.add_paragraph(style="List Number")
        p.add_run(h)
    doc.add_page_break()

    i = 0
    in_meta = False
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        if i == 0 and line == "---":
            in_meta = True
            i += 1
            continue
        if in_meta:
            if line == "---":
                in_meta = False
            i += 1
            continue
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph(style="No Spacing")
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.right_indent = Inches(0.18)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.add_run("\n".join(code_lines))
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if (
            line.startswith("|")
            and i + 1 < len(lines)
            and lines[i + 1].startswith("|")
            and "---" in lines[i + 1]
        ):
            rows = []
            rows.append([x.strip() for x in line.strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([x.strip() for x in lines[i].strip("|").split("|")])
                i += 1
            add_fixed_table(doc, rows)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:].strip())
        elif line.strip():
            p = doc.add_paragraph()
            p.add_run(line.strip())
        i += 1
    doc.core_properties.title = "GENESIS Complete System Design Specification"
    doc.core_properties.subject = (
        "Implementation-ready local-first generative agent-based modelling system specification"
    )
    doc.save(OUT)


if __name__ == "__main__":
    build_direct()
